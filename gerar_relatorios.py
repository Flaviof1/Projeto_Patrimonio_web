import os
import unicodedata
from datetime import datetime
import pandas as pd
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from identidade_visual import aplicar_marca_cabecalho_excel

from paths import DADOS_TRABALHO_DIR, caminho_relatorio, caminho_trabalho, garantir_pastas, localizar_arquivo
from dados_patrimoniais import carregar_dados_patrimoniais, enriquecer_com_dados_patrimoniais

ARQUIVO_TRABALHO = caminho_trabalho('conferencia_campo.xlsx')
ARQUIVO_RESULTADO = caminho_trabalho('conferencia_resultado.xlsx')
PREFIXO_RELATORIO = 'relatorio_patrimonial'


def obter_arquivo_base():
    """Obtém o arquivo de conferência principal"""
    resultado = localizar_arquivo(
        'conferencia_resultado.xlsx', DADOS_TRABALHO_DIR)
    if os.path.exists(resultado):
        return resultado

    trabalho = localizar_arquivo('conferencia_campo.xlsx', DADOS_TRABALHO_DIR)
    if os.path.exists(trabalho):
        return trabalho

    raise FileNotFoundError(
        f"Arquivo não encontrado: conferencia_resultado.xlsx ou conferencia_campo.xlsx"
    )


def obter_arquivo_localizar():
    """Mantido por compatibilidade; os valores vêm do CSV oficial de entrada."""
    return None


def gerar_nome_relatorio():
    garantir_pastas()
    data_hora = datetime.now().strftime('%Y-%m-%d_%H-%M')
    return caminho_relatorio(f'{PREFIXO_RELATORIO}_{data_hora}.xlsx')


def carregar_dados_conferencia(path=None):
    caminho = path or obter_arquivo_base()
    if not os.path.exists(caminho):
        raise FileNotFoundError(f"Arquivo não encontrado: {caminho}")
    return pd.read_excel(caminho).fillna('')


def carregar_dados_localizar(path=None):
    """Carrega o CSV oficial que contém os valores patrimoniais."""
    return carregar_dados_patrimoniais(path)


def validar_colunas(df):
    if 'STATUS' not in df.columns:
        raise ValueError('Coluna STATUS não encontrada.')


def normalizar_texto(valor):
    texto = str(valor).strip().lower()
    texto = unicodedata.normalize('NFKD', texto)
    return ''.join(ch for ch in texto if not unicodedata.combining(ch))


def coluna_texto(df, coluna):
    if coluna not in df.columns:
        return pd.Series([''] * len(df), index=df.index)
    return df[coluna].fillna('').astype(str)


def converter_para_numero(valor):
    if valor is None or valor == '' or valor == ' ' or valor == '-':
        return 0
    try:
        if isinstance(valor, str):
            valor_limpo = valor.strip().replace(',', '.')
            import re
            valor_limpo = re.sub(r'[^\d.\-]', '', valor_limpo)
            if valor_limpo == '':
                return 0
            return float(valor_limpo)
        if isinstance(valor, (int, float)):
            return float(valor)
        return 0
    except (ValueError, TypeError):
        return 0


def combinar_dados_com_valores(df_conferencia, df_valores):
    """Acrescenta valores por código e, na falta dele, por plaqueta."""
    return enriquecer_com_dados_patrimoniais(df_conferencia, df_valores)


def preparar_classificacao(df):
    validar_colunas(df)
    dados = df.copy()
    dados['STATUS'] = coluna_texto(dados, 'STATUS').str.upper().str.strip()

    observacao = coluna_texto(dados, 'OBSERVACAO').map(normalizar_texto)
    baixados = (dados['STATUS'] == 'BAIXA') | (
        (dados['STATUS'] == 'OK') & observacao.str.contains(
            'baixa', regex=False)
    )

    dados['CLASSIFICACAO_GERENCIAL'] = dados['STATUS']
    dados.loc[baixados, 'CLASSIFICACAO_GERENCIAL'] = 'BAIXADO'
    dados.loc[dados['STATUS'] == 'ALIEN',
              'CLASSIFICACAO_GERENCIAL'] = 'ALIENADO'
    dados.loc[(dados['STATUS'] == 'OK') & ~baixados,
              'CLASSIFICACAO_GERENCIAL'] = 'OK_CONFERIDO'
    return dados


def contar_pendencias_apos_2017(df):
    # Tentar encontrar coluna de data
    coluna_data = None
    for col in df.columns:
        if 'data' in col.lower() or 'aquisição' in col.lower():
            coluna_data = col
            break

    if coluna_data is None:
        return 0

    datas = pd.to_datetime(df[coluna_data], errors='coerce', dayfirst=True)
    mascara_ano_invalido = datas.dt.year < 1900
    datas.loc[mascara_ano_invalido] = pd.NaT

    pendentes = df[df['STATUS'].astype(str).str.upper() == 'PEND']
    return int(pendentes[datas.dt.year >= 2017].shape[0])


def criar_resumo_geral(df):
    foto_coluna = df.get('Foto', pd.Series([''] * len(df), index=df.index))
    fotos = foto_coluna.fillna('').astype(str).str.strip()

    ordem = [
        ('Total de bens', len(df)),
        ('OK conferidos', int(
            (df['CLASSIFICACAO_GERENCIAL'] == 'OK_CONFERIDO').sum())),
        ('Baixados', int((df['CLASSIFICACAO_GERENCIAL'] == 'BAIXADO').sum())),
        ('Alienados', int(
            (df['CLASSIFICACAO_GERENCIAL'] == 'ALIENADO').sum())),
        ('Pendentes', int((df['STATUS'] == 'PEND').sum())),
        ('Divergências', int((df['STATUS'] == 'DIVERG').sum())),
        ('Não localizados', int((df['STATUS'] == 'N_LOC').sum())),
        ('Sem plaqueta', int((df['STATUS'] == 'SEM_PLAQ').sum())),
        ('Não emplacados', int((df['STATUS'] == 'NAO_EMPLAC').sum())),
        ('Itens com foto anexada', int((fotos != '').sum())),
        ('Itens sem foto anexada', int((fotos == '').sum())),
    ]
    return pd.DataFrame(ordem, columns=['Indicador', 'Quantidade'])


def criar_resumo_por_setor(df):
    coluna_setor = None
    for col in df.columns:
        if 'setor' in col.lower():
            coluna_setor = col
            break

    if coluna_setor is None:
        return pd.DataFrame(columns=['Setor Atual'])

    resumo = pd.pivot_table(
        df,
        index=coluna_setor,
        columns='CLASSIFICACAO_GERENCIAL',
        values='STATUS',
        aggfunc='count',
        fill_value=0,
    ).reset_index()

    for coluna in ['OK_CONFERIDO', 'BAIXADO', 'PEND', 'DIVERG', 'N_LOC', 'SEM_PLAQ', 'NAO_EMPLAC']:
        if coluna not in resumo.columns:
            resumo[coluna] = 0

    for col in ['OK_CONFERIDO', 'BAIXADO', 'PEND', 'DIVERG', 'N_LOC', 'SEM_PLAQ', 'NAO_EMPLAC']:
        resumo[col] = pd.to_numeric(resumo[col], errors='coerce').fillna(0)

    resumo['TOTAL'] = resumo[
        ['OK_CONFERIDO', 'BAIXADO', 'PEND', 'DIVERG',
            'N_LOC', 'SEM_PLAQ', 'NAO_EMPLAC']
    ].sum(axis=1)

    resumo['% CONCLUIDO'] = 0.0

    for idx in resumo.index:
        total = resumo.loc[idx, 'TOTAL']
        if total > 0:
            concluidos = resumo.loc[idx, 'OK_CONFERIDO'] + \
                resumo.loc[idx, 'BAIXADO']
            resumo.loc[idx, '% CONCLUIDO'] = round(
                (concluidos / total) * 100, 1)

    resumo = resumo.rename(columns={coluna_setor: 'Setor Atual'})

    colunas = [
        'Setor Atual', 'TOTAL', 'OK_CONFERIDO', 'BAIXADO', 'PEND', 'DIVERG',
        'N_LOC', 'SEM_PLAQ', 'NAO_EMPLAC', '% CONCLUIDO'
    ]
    return resumo[colunas].sort_values(['PEND', 'DIVERG', 'N_LOC'], ascending=False)


def criar_resumo_divergentes_por_setor(df):
    coluna_setor = None
    for col in df.columns:
        if 'setor' in col.lower():
            coluna_setor = col
            break

    if coluna_setor is None or 'STATUS' not in df.columns:
        return pd.DataFrame(columns=['Setor Atual', 'Quantidade', 'Observação'])

    divergentes = df[df['STATUS'].astype(str).str.upper() == 'DIVERG'].copy()

    if divergentes.empty:
        return pd.DataFrame(columns=['Setor Atual', 'Quantidade', 'Observação'])

    resumo = divergentes.groupby(
        coluna_setor).size().reset_index(name='Quantidade')
    resumo = resumo.rename(columns={coluna_setor: 'Setor Atual'})

    obs_agrupadas = divergentes.groupby(coluna_setor)['OBSERVACAO'].apply(
        lambda x: '; '.join(x.dropna().unique())
    ).reset_index(name='Observações')
    obs_agrupadas = obs_agrupadas.rename(columns={coluna_setor: 'Setor Atual'})

    resumo = resumo.merge(obs_agrupadas, on='Setor Atual', how='left')
    return resumo.sort_values('Quantidade', ascending=False)


def criar_detalhamento_completo(df):
    # Mapear colunas dinamicamente
    mapeamento = {}

    colunas_interesse = {
        'Código': 'Código',
        'Plaqueta': 'Plaqueta',
        'Identificador Provisório': 'Identificador Provisório',
        'Descrição': 'Descrição',
        'Modelo': 'Modelo',
        'Setor Atual': 'Setor Atual',
        'Data': 'Data de Aquisição',
        'Valor Inicial': 'Valor Inicial (R$)',
        'Valor Atual': 'Valor Atual (R$)',
        'Vida Útil Inicial': 'Vida Útil Inicial (meses)',
        'Vida Útil Atual': 'Vida Útil Atual (meses)',
        'Valor Residual': 'Valor Residual (R$)',
        'Classificação': 'Classificação',
        'STATUS': 'Status Conferência',
        'ESTADO': 'Estado de Conservação',
        'OBSERVACAO': 'Observação',
        'Foto': 'Foto',
        'ALTERADO_EM': 'Última Alteração'
    }

    for col_origem, col_destino in colunas_interesse.items():
        if col_origem in df.columns:
            mapeamento[col_origem] = col_destino
        else:
            # Tentar encontrar coluna similar
            for col in df.columns:
                if col_origem.lower() in col.lower() or col.lower() in col_origem.lower():
                    mapeamento[col] = col_destino
                    break

    detalhes = pd.DataFrame()
    for col_origem, col_destino in mapeamento.items():
        if col_origem in df.columns:
            detalhes[col_destino] = df[col_origem]

    if 'Foto' in df.columns:
        detalhes['Possui Foto'] = df['Foto'].fillna('').astype(str).str.strip().map(
            lambda valor: 'Sim' if valor else 'Não'
        )

    # Converter colunas numéricas
    colunas_numericas = ['Valor Inicial (R$)', 'Valor Atual (R$)', 'Vida Útil Inicial (meses)',
                         'Vida Útil Atual (meses)', 'Valor Residual (R$)']

    for col in colunas_numericas:
        if col in detalhes.columns:
            detalhes[col] = detalhes[col].apply(converter_para_numero)

    # Calcular depreciação
    if 'Valor Inicial (R$)' in detalhes.columns and 'Vida Útil Atual (meses)' in detalhes.columns:
        def calcular_depreciacao(row):
            valor_inicial = converter_para_numero(row['Valor Inicial (R$)'])
            vida_util = converter_para_numero(row['Vida Útil Atual (meses)'])
            if vida_util > 0 and valor_inicial > 0:
                return round(valor_inicial / vida_util, 2)
            return ''

        detalhes['Depreciação Mensal (R$)'] = detalhes.apply(
            calcular_depreciacao, axis=1)

    if 'Valor Inicial (R$)' in detalhes.columns and 'Valor Atual (R$)' in detalhes.columns:
        def calcular_percentual_depreciacao(row):
            valor_inicial = converter_para_numero(row['Valor Inicial (R$)'])
            valor_atual = converter_para_numero(row['Valor Atual (R$)'])
            if valor_inicial > 0:
                return round(((valor_inicial - valor_atual) / valor_inicial * 100), 2)
            return ''

        detalhes['% Depreciação'] = detalhes.apply(
            calcular_percentual_depreciacao, axis=1)

    # Calcular idade
    coluna_data = None
    for col in detalhes.columns:
        if 'data' in col.lower() or 'aquisição' in col.lower():
            coluna_data = col
            break

    if coluna_data:
        def calcular_idade(data):
            if pd.isna(data) or data == '':
                return ''
            try:
                if isinstance(data, str):
                    data = pd.to_datetime(data, dayfirst=True, errors='coerce')
                if pd.isna(data):
                    return ''
                anos = (datetime.now() - data).days / 365.25
                return round(anos, 1)
            except:
                return ''

        detalhes['Idade (anos)'] = detalhes[coluna_data].apply(calcular_idade)

    return detalhes


def ajustar_planilha(writer, sheet_name):
    ws = writer.sheets[sheet_name]
    aplicar_marca_cabecalho_excel(ws)
    header_fill = PatternFill('solid', fgColor='1E3A8A')
    header_font = Font(color='FFFFFF', bold=True)
    border = Border(bottom=Side(style='thin', color='D8E0EC'))

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal='center')

    for column_cells in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            valor = '' if cell.value is None else str(cell.value)
            max_length = max(max_length, len(valor))
        ws.column_dimensions[column_letter].width = min(
            max(max_length + 2, 12), 45)

    for cabecalho in ws[1]:
        texto_cabecalho = str(cabecalho.value).lower()
        if 'valor' in texto_cabecalho or 'depreciação mensal' in texto_cabecalho:
            for coluna in ws.iter_cols(min_col=cabecalho.column, max_col=cabecalho.column, min_row=2):
                for celula in coluna:
                    if isinstance(celula.value, (int, float)):
                        celula.number_format = 'R$ #,##0.00'


def escrever_aba(writer, df, nome):
    if df is not None and not df.empty:
        df.to_excel(writer, sheet_name=nome, index=False)
        ajustar_planilha(writer, nome)
    else:
        df_vazio = pd.DataFrame(
            {'Mensagem': ['Nenhum dado disponível para esta aba']})
        df_vazio.to_excel(writer, sheet_name=nome, index=False)
        ajustar_planilha(writer, nome)


def exportar_relatorio_gerencial(df, path=None):
    caminho_saida = path or gerar_nome_relatorio()
    dados = preparar_classificacao(df)

    divergentes_sem_plaq = dados[
        (dados['STATUS'] == 'DIVERG') &
        (dados['OBSERVACAO'].astype(str).str.contains(
            'LOCALIZADO SEM PLAQUETA', case=False, na=False))
    ].copy()

    itens_localizados_sem_plaqueta = dados[
        coluna_texto(dados, 'Identificador Provisório').str.startswith(
            'SEM-ID-')
    ].copy()

    detalhamento_completo = criar_detalhamento_completo(dados)

    abas = {
        'Resumo_Geral': criar_resumo_geral(dados),
        'Resumo_por_Setor': criar_resumo_por_setor(dados),
        'Resumo_Divergentes': criar_resumo_divergentes_por_setor(dados),
        'Detalhamento_Completo': detalhamento_completo,
        'Baixados': dados[dados['CLASSIFICACAO_GERENCIAL'] == 'BAIXADO'],
        'OK_Conferidos': dados[dados['CLASSIFICACAO_GERENCIAL'] == 'OK_CONFERIDO'],
        'Pendencias': dados[dados['STATUS'] == 'PEND'],
        'Divergencias': dados[dados['STATUS'] == 'DIVERG'],
        'Divergentes_Sem_Plaqueta': divergentes_sem_plaq,
        'Itens_Localizados_Sem_Plaqueta': itens_localizados_sem_plaqueta,
        'Nao_Localizados': dados[dados['STATUS'] == 'N_LOC'],
        'Sem_Plaqueta': dados[dados['STATUS'] == 'SEM_PLAQ'],
        'Nao_Emplacados': dados[dados['STATUS'] == 'NAO_EMPLAC'],
    }

    with pd.ExcelWriter(caminho_saida, engine='openpyxl') as writer:
        for nome, dados_aba in abas.items():
            escrever_aba(writer, dados_aba, nome)

    return {
        'arquivo': caminho_saida,
        'total': len(dados),
        'baixados': len(abas['Baixados']),
        'ok_conferidos': len(abas['OK_Conferidos']),
        'pendencias': len(abas['Pendencias']),
        'divergencias': len(abas['Divergencias']),
        'divergentes_sem_plaq': len(divergentes_sem_plaq),
        'itens_localizados_sem_plaqueta': len(itens_localizados_sem_plaqueta),
        'nao_localizados': len(abas['Nao_Localizados']),
        'sem_plaqueta': len(abas['Sem_Plaqueta']),
        'nao_emplacados': len(abas['Nao_Emplacados']),
        'total_abas': len(abas),
    }
