# utils.py
# Funções utilitárias sem dependência de tkinter (formatação, exportação de relatórios)

import os
import re
import unicodedata
from copy import copy
from datetime import datetime
import pandas as pd
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from identidade_visual import aplicar_marca_cabecalho_excel, desenhar_marca_cabecalho_pdf
from paths import caminho_relatorio, garantir_pastas
from constantes import STATUS_OPCOES, PREFIXO_RELATORIO_BUSCA

CORES_STATUS_RELATORIO = {
    'PEND': 'C00000',
    'N_LOC': 'C00000',
    'BAIXA': '0070C0',
    'ALIEN': '0070C0',
}


def formatar_valor_reais(valor):
    """Formata um valor numérico para o formato R$ 1.234,56"""
    if valor is None or pd.isna(valor) or str(valor).strip() == '':
        return ''
    try:
        numero = float(valor)
        texto = f'{numero:,.2f}'
        return f"R$ {texto.replace(',', 'X').replace('.', ',').replace('X', '.')}"
    except (TypeError, ValueError):
        return str(valor)


def converter_valor_atual(valor):
    """Converte valores numéricos ou monetários para float, usando zero se inválido."""
    if valor is None or pd.isna(valor):
        return 0.0
    if isinstance(valor, (int, float)):
        return float(valor)

    texto_original = str(valor).strip()
    texto = re.sub(r'[^\d,.+\-]', '', texto_original)
    if not texto:
        return 0.0

    negativo_parenteses = texto_original.startswith(
        '(') and texto_original.endswith(')')
    if ',' in texto and '.' in texto:
        if texto.rfind(',') > texto.rfind('.'):
            texto = texto.replace('.', '').replace(',', '.')
        else:
            texto = texto.replace(',', '')
    elif ',' in texto:
        partes = texto.split(',')
        if len(partes) > 2:
            if len(partes[-1]) <= 2:
                texto = ''.join(partes[:-1]) + '.' + partes[-1]
            else:
                texto = ''.join(partes)
        elif len(partes[-1]) == 3:
            texto = ''.join(partes)
        else:
            texto = '.'.join(partes)
    elif texto.count('.') > 1:
        partes = texto.split('.')
        if len(partes[-1]) <= 2:
            texto = ''.join(partes[:-1]) + '.' + partes[-1]
        else:
            texto = ''.join(partes)
    elif '.' in texto and len(texto.rsplit('.', 1)[-1]) == 3:
        texto = texto.replace('.', '')

    try:
        numero = float(texto)
        return -abs(numero) if negativo_parenteses else numero
    except ValueError:
        return 0.0


def valores_atuais_numericos(dados):
    """Retorna os valores atuais convertidos para número, mantendo o índice."""
    if 'Valor Atual' not in dados.columns:
        return pd.Series(0.0, index=dados.index)
    return dados['Valor Atual'].map(converter_valor_atual).astype(float)


def somar_valor_atual(dados):
    """Soma o valor atual dos itens informados."""
    return float(valores_atuais_numericos(dados).sum())


def criar_resumo_por_categoria(dados, coluna, nome_coluna, valor_vazio):
    categorias = dados[coluna].fillna('').astype(str).str.strip()
    if coluna == 'STATUS':
        categorias = categorias.str.upper()
    categorias = categorias.replace('', valor_vazio)
    agrupamento = pd.DataFrame({
        nome_coluna: categorias,
        '_valor_atual': valores_atuais_numericos(dados),
    })
    return (
        agrupamento.groupby(nome_coluna, as_index=False, sort=False)
        .agg(
            Quantidade=('_valor_atual', 'size'),
            **{'Valor Atual Total (R$)': ('_valor_atual', 'sum')},
        )
        .sort_values('Quantidade', ascending=False)
        .reset_index(drop=True)
    )


def gerar_nome_relatorio_busca():
    """Gera nome único para relatório de busca com timestamp"""
    garantir_pastas()
    data_hora = datetime.now().strftime('%Y-%m-%d_%H-%M')
    return caminho_relatorio(f'{PREFIXO_RELATORIO_BUSCA}_{data_hora}.xlsx')


def gerar_nome_pdf_busca():
    """Gera nome para o PDF dos itens filtrados na busca geral."""
    garantir_pastas()
    data_hora = datetime.now().strftime('%Y-%m-%d_%H-%M')
    return caminho_relatorio(f'{PREFIXO_RELATORIO_BUSCA}_{data_hora}.pdf')


def _ordenar_por_descricao(dados):
    if 'Descrição' not in dados.columns:
        return dados.copy()

    def chave_alfabetica(valor):
        if pd.isna(valor):
            return pd.NA
        texto = unicodedata.normalize('NFKD', str(valor).casefold())
        return ''.join(
            caractere for caractere in texto
            if not unicodedata.combining(caractere)
        )

    return dados.sort_values(
        by='Descrição',
        key=lambda coluna: coluna.map(chave_alfabetica),
        kind='stable',
        na_position='last',
    )


def exportar_relatorio_busca_pdf(dados, caminho_saida=None, filtros=None):
    """Exporta os itens filtrados da busca geral para PDF em paisagem."""
    if dados.empty:
        raise ValueError('Não há dados para exportar.')

    dados_ordenados = _ordenar_por_descricao(dados)
    caminho = caminho_saida or gerar_nome_pdf_busca()
    colunas_exportacao = [
        ('Plaqueta', 'Número da Plaqueta'),
        ('Descrição', 'Descrição'),
        ('Setor Atual', 'Setor Atual'),
        ('STATUS', 'Status'),
        ('ESTADO', 'Estado'),
        ('OBSERVACAO', 'Observação'),
        ('Data', 'Data do Item'),
        ('Valor Atual', 'Valor Atual'),
    ]

    estilos = getSampleStyleSheet()
    estilo_celula = estilos['BodyText'].clone('celula_pdf')
    estilo_celula.fontSize = 7
    estilo_celula.leading = 8
    estilo_cabecalho = estilos['BodyText'].clone('cabecalho_pdf')
    estilo_cabecalho.fontSize = 7
    estilo_cabecalho.leading = 8
    estilo_cabecalho.textColor = colors.white
    estilo_cabecalho.fontName = 'Helvetica-Bold'

    def texto(valor):
        if pd.isna(valor):
            return ''
        return str(valor).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

    cabecalho = [Paragraph(texto(rotulo), estilo_cabecalho)
                 for _, rotulo in colunas_exportacao]
    linhas = [cabecalho]
    for _, linha in dados_ordenados.iterrows():
        celulas = []
        for coluna, _ in colunas_exportacao:
            valor = linha.get(coluna, '')
            if coluna == 'Valor Atual':
                valor = formatar_valor_reais(valor)
            conteudo = texto(valor)
            if coluna == 'STATUS' and pd.notna(valor):
                cor = CORES_STATUS_RELATORIO.get(str(valor).strip().upper())
                if cor:
                    conteudo = f'<font color="#{cor}"><b>{conteudo}</b></font>'
            celulas.append(Paragraph(conteudo, estilo_celula))
        linhas.append(celulas)

    documento = SimpleDocTemplate(
        caminho,
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=16 * mm,
        bottomMargin=10 * mm,
        title='Relatório da Busca Geral',
    )
    titulo = estilos['Title']
    titulo.fontSize = 14
    titulo.leading = 16
    estilo_filtros = estilos['Normal'].clone('filtros_pdf')
    estilo_filtros.fontSize = 8
    estilo_filtros.leading = 10
    largura_util = landscape(A4)[0] - 20 * mm
    larguras = [
        largura_util * 0.12,
        largura_util * 0.21,
        largura_util * 0.14,
        largura_util * 0.09,
        largura_util * 0.10,
        largura_util * 0.15,
        largura_util * 0.10,
        largura_util * 0.09,
    ]
    tabela = Table(linhas, colWidths=larguras, repeatRows=1)
    tabela.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('GRID', (0, 0), (-1, -1), 0.35, colors.black),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('BACKGROUND', (0, 1), (-1, -1), colors.white),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1),
         [colors.white, colors.HexColor('#F2F2F2')]),
        ('LEFTPADDING', (0, 0), (-1, -1), 3),
        ('RIGHTPADDING', (0, 0), (-1, -1), 3),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]))
    filtros = filtros or {}
    termo = filtros.get('termo') or 'todos'
    setores = ', '.join(filtros.get('setores') or []) or 'todos'
    status = ', '.join(filtros.get('status') or []) or 'todos'
    fotos = ', '.join(filtros.get('foto') or []) or 'todos'
    plaqueta = filtros.get('plaqueta') or 'todas'
    data_inicial = filtros.get('data_inicial') or 'não informada'
    data_final = filtros.get('data_final') or 'não informada'
    descricao_filtros = (
        f'<b>Filtros aplicados:</b> termo: {texto(termo)} | '
        f'setor(es): {texto(setores)} | status: {texto(status)} | '
        f'plaqueta: {texto(plaqueta)} | foto: {texto(fotos)} | '
        f'período: {texto(data_inicial)} até {texto(data_final)}'
    )

    documento.build([
        Paragraph('Relatório da Busca Geral', titulo),
        Paragraph(descricao_filtros, estilo_filtros),
        Paragraph(
            f'{len(dados)} item(ns) filtrado(s) - gerado em {datetime.now():%d/%m/%Y %H:%M}', estilos['Normal']),
        Paragraph(
            f'<b>Valor Atual total: {texto(formatar_valor_reais(somar_valor_atual(dados)))}</b>',
            estilos['Normal']),
        Spacer(1, 5 * mm),
        tabela,
    ], onFirstPage=desenhar_marca_cabecalho_pdf,
        onLaterPages=desenhar_marca_cabecalho_pdf)
    return caminho


def ajustar_planilha_relatorio(writer, sheet_name):
    """Aplica formatação (cabeçalho, largura, bordas) a uma aba do relatório"""
    ws = writer.sheets[sheet_name]
    aplicar_marca_cabecalho_excel(ws)
    header_fill = PatternFill('solid', fgColor='1E3A8A')
    header_font = Font(color='FFFFFF', bold=True)
    border = Border(bottom=Side(style='thin', color='D8E0EC'))

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal='center')

        status = str(cell.value or '').strip().upper()
        cor = CORES_STATUS_RELATORIO.get(status)
        if cor:
            cell.fill = PatternFill(
                'solid', fgColor='FCE4D6' if cor == 'C00000' else 'DDEBF7')
            cell.font = Font(color=cor, bold=True)

    for coluna in ws[1]:
        if str(coluna.value or '').strip().upper() != 'STATUS':
            continue
        for celula in ws.iter_cols(
                min_col=coluna.column, max_col=coluna.column,
                min_row=2, max_row=ws.max_row):
            for item in celula:
                cor = CORES_STATUS_RELATORIO.get(
                    str(item.value or '').strip().upper())
                if cor:
                    fonte = copy(item.font)
                    fonte.color = cor
                    fonte.bold = True
                    item.font = fonte

    for column_cells in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            valor = '' if cell.value is None else str(cell.value)
            max_length = max(max_length, len(valor))
        ws.column_dimensions[column_letter].width = min(
            max(max_length + 2, 12), 45)

    for cabecalho in ws[1]:
        if 'valor' in str(cabecalho.value).lower():
            for coluna in ws.iter_cols(min_col=cabecalho.column, max_col=cabecalho.column, min_row=2):
                for celula in coluna:
                    if isinstance(celula.value, (int, float)):
                        celula.number_format = 'R$ #,##0.00'


def escrever_aba_relatorio(writer, dados, nome):
    """Escreve um DataFrame em uma aba e aplica formatação"""
    dados.to_excel(writer, sheet_name=nome, index=False)
    ajustar_planilha_relatorio(writer, nome)


def criar_resumo_setor_status(dados):
    """Cria tabela cruzada de setor x status"""
    if 'Setor Atual' not in dados.columns or 'STATUS' not in dados.columns:
        return pd.DataFrame()

    resumo = pd.pivot_table(
        dados,
        index='Setor Atual',
        columns='STATUS',
        values=dados.columns[0],
        aggfunc='count',
        fill_value=0,
    ).reset_index()

    resumo['TOTAL'] = resumo.drop(columns=['Setor Atual']).sum(axis=1)
    colunas_status = [
        coluna for coluna in STATUS_OPCOES if coluna in resumo.columns]
    outras_colunas = [
        coluna for coluna in resumo.columns
        if coluna not in ['Setor Atual', 'TOTAL'] + colunas_status
    ]
    return resumo[['Setor Atual', 'TOTAL'] + colunas_status + outras_colunas]


def criar_resumo_valor_setor_status(dados):
    """Cria tabela cruzada de valores atuais por setor e status."""
    if 'Setor Atual' not in dados.columns or 'STATUS' not in dados.columns:
        return pd.DataFrame()

    agrupamento = pd.DataFrame({
        'Setor Atual': (
            dados['Setor Atual'].fillna('').astype(str).str.strip()
            .replace('', 'Sem setor')
        ),
        'STATUS': (
            dados['STATUS'].fillna('').astype(str).str.strip().str.upper()
            .replace('', 'Sem status')
        ),
        '_valor_atual': valores_atuais_numericos(dados),
    })
    resumo = pd.pivot_table(
        agrupamento,
        index='Setor Atual',
        columns='STATUS',
        values='_valor_atual',
        aggfunc='sum',
        fill_value=0,
    ).reset_index()
    resumo.columns.name = None

    colunas_status = [
        coluna for coluna in resumo.columns if coluna != 'Setor Atual']
    colunas_status = (
        [coluna for coluna in STATUS_OPCOES if coluna in colunas_status] +
        [coluna for coluna in colunas_status if coluna not in STATUS_OPCOES]
    )
    resumo['Valor Atual Total (R$)'] = resumo[colunas_status].sum(axis=1)
    return resumo[['Setor Atual'] + colunas_status + ['Valor Atual Total (R$)']]


def exportar_relatorio_busca(dados, caminho_saida=None):
    """Exporta os dados da busca para um arquivo Excel com várias abas"""
    caminho = caminho_saida or gerar_nome_relatorio_busca()
    dados_exportacao = dados.copy()

    abas = {'Itens_Filtrados': _ordenar_por_descricao(dados_exportacao)}

    if 'Setor Atual' in dados_exportacao.columns:
        abas['Resumo_por_Setor'] = criar_resumo_por_categoria(
            dados_exportacao, 'Setor Atual', 'Setor Atual', 'Sem setor')

    if 'STATUS' in dados_exportacao.columns:
        abas['Resumo_por_Status'] = criar_resumo_por_categoria(
            dados_exportacao, 'STATUS', 'STATUS', 'Sem status')

    resumo_setor_status = criar_resumo_setor_status(dados_exportacao)
    if not resumo_setor_status.empty:
        abas['Setor_x_Status'] = resumo_setor_status

    resumo_valor_setor_status = criar_resumo_valor_setor_status(
        dados_exportacao)
    if not resumo_valor_setor_status.empty:
        abas['Valor_Setor_x_Status'] = resumo_valor_setor_status

    with pd.ExcelWriter(caminho, engine='openpyxl') as writer:
        for nome, dados_aba in abas.items():
            escrever_aba_relatorio(writer, dados_aba, nome)
            if nome == 'Itens_Filtrados':
                planilha = writer.sheets[nome]
                linha_total = planilha.max_row + 1
                rotulo_total = planilha.cell(
                    row=linha_total, column=1, value='TOTAL DOS RESULTADOS')
                rotulo_total.font = Font(bold=True)
                rotulo_total.fill = PatternFill('solid', fgColor='DDEBF7')

                colunas = {
                    str(celula.value): celula.column
                    for celula in planilha[1]
                }
                coluna_valor = colunas.get('Valor Atual')
                if coluna_valor:
                    celula_total = planilha.cell(
                        row=linha_total,
                        column=coluna_valor,
                        value=somar_valor_atual(dados_exportacao),
                    )
                    celula_total.number_format = 'R$ #,##0.00'
                    celula_total.font = Font(bold=True)
                    celula_total.fill = PatternFill(
                        'solid', fgColor='DDEBF7')

    return caminho


def exportar_dados_excel(dados, nome_base, titulo_aba="Itens"):
    """
    Exporta um DataFrame para Excel com formatação profissional.

    Args:
        dados (pd.DataFrame): dados a exportar
        nome_base (str): prefixo do nome do arquivo (ex: 'conferencia_setor_X')
        titulo_aba (str): nome da aba no Excel

    Returns:
        str: caminho do arquivo gerado
    """
    from paths import caminho_relatorio, garantir_pastas
    from datetime import datetime

    if dados.empty:
        raise ValueError("Não há dados para exportar.")

    garantir_pastas()
    data_hora = datetime.now().strftime('%Y-%m-%d_%H-%M')
    caminho = caminho_relatorio(f'{nome_base}_{data_hora}.xlsx')

    with pd.ExcelWriter(caminho, engine='openpyxl') as writer:
        dados.to_excel(writer, sheet_name=titulo_aba, index=False)
        ajustar_planilha_relatorio(writer, titulo_aba)

    return caminho
