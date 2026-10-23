import os
import shutil
import tempfile
from pathlib import Path
from io import BytesIO

import pytest


_TEST_DATA = tempfile.TemporaryDirectory(prefix='patrimonio_web_test_')
os.environ['PATRIMONIO_WEB_DATA_DIR'] = _TEST_DATA.name

import pandas as pd
import zipfile
from PIL import Image
from openpyxl import load_workbook

import app as app_module
from app import app
from paths import (
    BACKUPS_DIR,
    DATA_DIR,
    DADOS_ENTRADA_DIR,
    DADOS_TRABALHO_DIR,
    FOTOS_DIR,
)
from preparar_dados import normalizar_plaqueta_numerica

CSV_VALIDO = (
    'Código;Plaqueta;Descrição;Modelo;Setor Atual;Data\n'
    '001;000123;Mesa;M1;Setor A;01/02/2024\n'
    '002;000124;Cadeira;C2;Setor B;15/03/2023\n'
)


@pytest.fixture(autouse=True)
def limpar_dados():
    pasta_real = Path(__file__).resolve().parents[1] / 'dados_web'
    assert DATA_DIR.resolve() != pasta_real.resolve(), (
        'Os testes devem usar uma pasta temporária, nunca dados_web do projeto.'
    )
    for pasta in ('dados_entrada', 'dados_trabalho', 'relatorios', 'backups'):
        caminho = DATA_DIR / pasta
        if caminho.exists():
            shutil.rmtree(caminho)
    yield


@pytest.fixture
def client():
    with app.test_client() as cliente:
        cliente.get('/')
        with cliente.session_transaction() as sessao:
            sessao['csrf_token'] = 'token-de-teste'
        yield cliente


def post(client, url, **kwargs):
    headers = kwargs.pop('headers', {})
    headers['X-CSRF-Token'] = 'token-de-teste'
    return client.post(url, headers=headers, **kwargs)


def enviar_csv(client, conteudo=CSV_VALIDO, confirmar=None):
    data = {'arquivo': (BytesIO(conteudo.encode('latin1')), 'cadastro.csv')}
    if confirmar:
        data['confirmar_substituicao'] = 'sim'
    return post(client, '/api/importar/csv', data=data, content_type='multipart/form-data')


def enviar_auxiliar(client, tipo, nome, conteudo):
    return post(
        client,
        '/api/importar/auxiliar',
        data={
            'tipo': tipo,
            'arquivo': (BytesIO(conteudo), nome),
        },
        content_type='multipart/form-data',
    )


def enviar_resultado(client, workbook, confirmar=None):
    data = {'arquivo': (BytesIO(workbook), 'conferencia_resultado.xlsx')}
    if confirmar:
        data['confirmado'] = 'sim'
    return post(
        client,
        '/api/importar/resultado',
        data=data,
        content_type='multipart/form-data',
    )


def criar_resultado_excel(linhas):
    arquivo = BytesIO()
    pd.DataFrame(linhas).to_excel(arquivo, index=False)
    return arquivo.getvalue()


def preparar_base(client):
    resposta = enviar_csv(client)
    assert resposta.status_code == 200, resposta.get_json()
    return post(client, '/api/preparar', data={})


def test_aplicacao_inicia_sem_abrir_dados_do_desktop(client):
    resposta = client.get('/api/status')

    assert resposta.status_code == 200
    assert resposta.get_json()['estado'] == 'sem_dados'
    assert not (
        DATA_DIR / 'dados_trabalho'
    ).joinpath('conferencia_resultado.xlsx').exists()


def test_modo_mysql_usa_a_mesma_camada_de_leitura_e_gravacao(client, monkeypatch):
    base = pd.DataFrame([{
        'Código': '001',
        'Plaqueta': '000123',
        'Descrição': 'Mesa',
        'Setor Atual': 'Setor A',
        'STATUS': 'PEND',
        'ESTADO': 'EM USO',
        'OBSERVACAO': '',
        'Foto': '',
    }])
    gravado = {}

    class ArmazenamentoTeste:
        @classmethod
        def from_environment(cls):
            return cls()

        def load(self):
            return base.copy()

        def replace_all(self, dados):
            gravado['dados'] = dados.copy()

    monkeypatch.setattr(app_module, 'MySQLStore', ArmazenamentoTeste)
    monkeypatch.setenv('PATRIMONIO_STORAGE', 'mysql')

    status = client.get('/api/status').get_json()
    assert status['arquivo'] == 'MySQL'
    assert status['total'] == 1

    item = client.get('/api/conferencia?setor=Setor%20A').get_json()['itens'][0]
    resposta = post(
        client,
        f"/api/itens/{item['id']}/editar",
        json={'status': 'OK', 'estado': 'DESUSO', 'observacao': 'Gravado no MySQL'},
    )

    assert resposta.status_code == 200
    assert len(gravado['dados']) == 1
    assert gravado['dados'].loc[item['id'], 'OBSERVACAO'] == 'Gravado no MySQL'


def test_normalizacao_plaqueta_nao_acrescenta_digito_apos_leitura_excel():
    assert normalizar_plaqueta_numerica(123.0) == '000123'
    assert normalizar_plaqueta_numerica('000123') == '000123'


def test_importacao_csv_valida_e_preserva_arquivo_ao_receber_csv_invalido(client):
    resposta = enviar_csv(client)
    assert resposta.status_code == 200
    assert resposta.get_json()['registros'] == 2
    arquivo = DADOS_ENTRADA_DIR / 'Localizar_Bens.csv'
    original = arquivo.read_bytes()

    invalido = enviar_csv(client, 'Sem as colunas obrigatórias\n1;2;3\n', confirmar=True)

    assert invalido.status_code == 400
    assert arquivo.read_bytes() == original


def test_importa_baixados_e_preparacao_marca_status_por_plaqueta(client):
    assert enviar_csv(client).status_code == 200
    planilha = BytesIO()
    pd.DataFrame({'Plaqueta': [123, 999]}).to_excel(planilha, index=False)
    importacao = enviar_auxiliar(
        client, 'baixados', 'Bens_Baixados.xlsx', planilha.getvalue()
    )
    assert importacao.status_code == 200, importacao.get_json()

    preparar = post(client, '/api/preparar', data={})

    assert preparar.status_code == 200, preparar.get_json()
    resumo = preparar.get_json()['baixados']
    assert resumo == {
        'na_planilha': 2,
        'encontrados_no_cadastro': 1,
        'marcados_baixa': 1,
        'prevaleceram_alienados': 0,
        'nao_encontrados': 1,
    }
    itens = client.get('/api/conferencia?setor=Setor%20A').get_json()['itens']
    baixado = next(item for item in itens if item['Plaqueta'] == '000123')
    assert baixado['STATUS'] == 'BAIXA'
    assert baixado['ESTADO'] == 'NAO_APLICAVEL'


def test_importacao_baixados_rejeita_planilha_sem_coluna_plaqueta(client):
    planilha = BytesIO()
    pd.DataFrame({'Codigo': ['001']}).to_excel(planilha, index=False)

    resposta = enviar_auxiliar(
        client, 'baixados', 'Bens_Baixados.xlsx', planilha.getvalue()
    )

    assert resposta.status_code == 400
    assert 'Plaqueta' in resposta.get_json()['error']


def test_interface_mantem_formularios_de_fontes_sem_ods_ou_importacao_zip(client):
    html = client.get('/').get_data(as_text=True)

    assert 'alienados-form' in html
    assert 'baixados-form' in html
    assert 'id="sector-form"' in html
    assert 'type="submit">Iniciar</button>' in html
    assert 'id="clear-search"' in html
    assert 'Limpar filtros' in html
    assert 'Comparar bens baixados' not in html
    assert 'Migração do desktop' not in html
    assert 'result-form' not in html
    assert 'photos-form' not in html
    assert 'sem-plaqueta-form' not in html
    assert post(
        client,
        '/api/importar/fotos',
        data={},
        content_type='multipart/form-data',
    ).status_code == 404


def test_importacao_resultado_migra_edicoes_e_preserva_referencias_de_fotos(client):
    planilha = criar_resultado_excel([{
        'Código': '001',
        'Plaqueta': '000123',
        'Descrição': 'Mesa conferida',
        'Setor Atual': 'Setor A',
        'STATUS': 'OK',
        'ESTADO': 'EM USO',
        'OBSERVACAO': 'Conferência migrada',
        'Foto': 'foto-desktop.jpg',
        'ALTERADO_EM': '20/10/2026 10:00:00',
    }])

    resposta = enviar_resultado(client, planilha)

    assert resposta.status_code == 200, resposta.get_json()
    assert resposta.get_json()['bens'] == 1
    assert resposta.get_json()['setores'] == 1
    assert resposta.get_json()['referencias_foto_preservadas'] == 1
    item = client.get('/api/conferencia?setor=Setor%20A').get_json()['itens'][0]
    assert item['STATUS'] == 'OK'
    assert item['ESTADO'] == 'EM USO'
    assert item['OBSERVACAO'] == 'Conferência migrada'
    assert item['ALTERADO_EM'] == '20/10/2026 10:00:00'
    assert item['Foto'] == 'foto-desktop.jpg'
    assert not (DADOS_ENTRADA_DIR / 'conferencia_resultado.xlsx').exists()


def test_migracao_resultado_exige_confirmacao_faz_backup_e_preserva_base_se_cancelada(client):
    assert preparar_base(client).status_code == 200
    planilha = criar_resultado_excel([{
        'Código': '901',
        'Plaqueta': '009999',
        'Descrição': 'Bem migrado',
        'Setor Atual': 'Setor Z',
        'STATUS': 'DIVERG',
        'ESTADO': 'DESUSO',
        'OBSERVACAO': 'Migrado',
    }])

    precisa_confirmar = enviar_resultado(client, planilha)
    assert precisa_confirmar.status_code == 409
    assert precisa_confirmar.get_json()['requires_confirmation'] is True
    assert precisa_confirmar.get_json()['preview']['bens'] == 1
    assert client.get('/api/conferencia?setor=Setor%20A').get_json()['total'] == 1

    migrado = enviar_resultado(client, planilha, confirmar=True)
    assert migrado.status_code == 200, migrado.get_json()
    assert client.get('/api/conferencia?setor=Setor%20Z').get_json()['total'] == 1
    assert list(BACKUPS_DIR.glob('backup_web_*.zip'))


@pytest.mark.parametrize(
    'linhas',
    [
        [{'Descrição': 'Sem colunas obrigatórias'}],
        [{
            'Plaqueta': '000123',
            'Descrição': 'Status desconhecido',
            'Setor Atual': 'Setor A',
            'STATUS': 'DESCONHECIDO',
        }],
    ],
)
def test_planilha_resultado_invalida_nao_substitui_base(linhas, client):
    assert preparar_base(client).status_code == 200
    resposta = enviar_resultado(client, criar_resultado_excel(linhas))

    assert resposta.status_code == 400
    assert client.get('/api/conferencia?setor=Setor%20A').get_json()['total'] == 1


def test_preparacao_setores_conferencia_e_regras_de_estado(client):
    preparar = preparar_base(client)
    assert preparar.status_code == 200, preparar.get_json()
    assert preparar.get_json()['total'] == 2

    setores = client.get('/api/setores').get_json()['setores']
    assert setores == ['Setor A', 'Setor B']

    conferencia = client.get('/api/conferencia?setor=Setor%20A').get_json()
    assert conferencia['total'] == 1
    item = conferencia['itens'][0]
    assert item['STATUS'] == 'PEND'

    edicao = post(
        client,
        f"/api/itens/{item['id']}/editar",
        json={'status': 'N_LOC', 'estado': 'EM USO', 'observacao': 'Não localizado'},
    )
    assert edicao.status_code == 200, edicao.get_json()
    atualizado = client.get('/api/conferencia?setor=Setor%20A').get_json()['itens'][0]
    assert atualizado['STATUS'] == 'N_LOC'
    assert atualizado['ESTADO'] == 'NAO_APLICAVEL'
    assert atualizado['OBSERVACAO'] == 'Não localizado'


def test_edicao_inclusao_busca_e_repreparacao_preservam_conferencia(client):
    assert preparar_base(client).status_code == 200
    item = client.get('/api/conferencia?setor=Setor%20A').get_json()['itens'][0]

    incluir = post(
        client,
        '/api/conferencia/incluir',
        json={
            'setor': 'Setor A',
            'descricao': 'Bem localizado sem plaqueta',
            'quantidade': 2,
            'status': 'DIVERG',
            'estado': 'EM USO',
        },
    )
    assert incluir.status_code == 200, incluir.get_json()

    busca = client.get('/api/busca?termo=plaqueta&setor=Setor+A&status=DIVERG')
    assert busca.status_code == 200
    assert busca.get_json()['total'] == 2

    editar = post(
        client,
        f"/api/itens/{item['id']}/editar",
        json={'status': 'OK', 'estado': 'EM USO', 'observacao': 'Conferido'},
    )
    assert editar.status_code == 200
    segunda_preparacao = post(client, '/api/preparar', data={'confirmado': 'sim'})
    assert segunda_preparacao.status_code == 200, segunda_preparacao.get_json()

    conferidos = client.get('/api/conferencia?setor=Setor%20A').get_json()['itens']
    item_original = next(
        (linha for linha in conferidos if str(linha['Plaqueta']) == '000123'),
        None,
    )
    assert item_original is not None, [
        (linha.get('Plaqueta'), linha.get('Setor Atual'), linha.get('STATUS'))
        for linha in conferidos
    ]
    assert item_original['STATUS'] == 'OK'
    assert item_original['OBSERVACAO'] == 'Conferido'


def test_importa_fontes_auxiliares_e_preserva_regras_de_preparacao(client):
    assert enviar_csv(client).status_code == 200

    xlsx = BytesIO()
    pd.DataFrame([{
        'Plaqueta': '000123',
        'Bem / Descrição': 'Mesa alienada',
        'Data': '10/01/2024',
    }]).to_excel(xlsx, index=False)
    resposta_xlsx = enviar_auxiliar(
        client, 'alienados', 'Bens_Alienados.xlsx', xlsx.getvalue()
    )
    assert resposta_xlsx.status_code == 200, resposta_xlsx.get_json()

    baixados_xlsx = BytesIO()
    pd.DataFrame({'Plaqueta': ['000123', '000124']}).to_excel(
        baixados_xlsx, index=False
    )
    resposta_baixados = enviar_auxiliar(
        client, 'baixados', 'Bens_Baixados.xlsx', baixados_xlsx.getvalue()
    )
    assert resposta_baixados.status_code == 200, resposta_baixados.get_json()

    ods = BytesIO()
    xml = '''<?xml version="1.0" encoding="UTF-8"?>
    <office:document-content
      xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
      xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0"
      xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">
      <office:body><office:spreadsheet><table:table>
        <table:table-row>
          <table:table-cell><text:p>Plaqueta</text:p></table:table-cell>
          <table:table-cell><text:p>Descrição</text:p></table:table-cell>
          <table:table-cell><text:p>Setor Atual</text:p></table:table-cell>
          <table:table-cell><text:p>Estado</text:p></table:table-cell>
          <table:table-cell><text:p>Observação</text:p></table:table-cell>
        </table:table-row>
        <table:table-row>
          <table:table-cell><text:p></text:p></table:table-cell>
          <table:table-cell><text:p>Bem achado</text:p></table:table-cell>
          <table:table-cell><text:p>Setor B</text:p></table:table-cell>
          <table:table-cell><text:p>DESUSO</text:p></table:table-cell>
          <table:table-cell><text:p>Sem plaqueta oficial</text:p></table:table-cell>
        </table:table-row>
      </table:table></office:spreadsheet></office:body>
    </office:document-content>'''.encode('utf-8')
    with zipfile.ZipFile(ods, 'w', zipfile.ZIP_DEFLATED) as arquivo_zip:
        arquivo_zip.writestr('content.xml', xml)
    resposta_ods = enviar_auxiliar(
        client, 'sem_plaqueta', 'itens_inclusos_sem_plaquetas.ods', ods.getvalue()
    )
    assert resposta_ods.status_code == 200, resposta_ods.get_json()

    preparar = post(client, '/api/preparar', data={})
    assert preparar.status_code == 200, preparar.get_json()
    assert preparar.get_json()['avisos'] == [
        '1 plaqueta(s) da lista de baixados também constam como alienadas; o status ALIEN foi mantido.'
    ]
    assert preparar.get_json()['baixados'] == {
        'na_planilha': 2,
        'encontrados_no_cadastro': 2,
        'marcados_baixa': 1,
        'prevaleceram_alienados': 1,
        'nao_encontrados': 0,
    }

    alienados = client.get('/api/conferencia?setor=Alienados').get_json()['itens']
    assert len(alienados) == 1
    assert alienados[0]['STATUS'] == 'ALIEN'
    achados = client.get('/api/conferencia?setor=Setor%20B').get_json()['itens']
    item_baixado = next(item for item in achados if item['Plaqueta'] == '000124')
    assert item_baixado['STATUS'] == 'BAIXA'
    assert any(item['Descrição'] == 'Bem achado' for item in achados)


def test_anexa_remove_e_filtra_foto_de_bem(client):
    assert preparar_base(client).status_code == 200
    assert client.get('/api/busca?foto=Sem+foto').get_json()['total'] == 2
    assert client.get('/api/busca?foto=Com+foto').get_json()['total'] == 0


    foto = BytesIO()
    Image.new('RGB', (1, 1), color=(20, 40, 60)).save(foto, format='PNG')
    resposta = post(
        client,
        '/api/itens/0/foto',
        data={
            'arquivo': (BytesIO(foto.getvalue()), 'foto.png'),
            'status': 'OK',
            'estado': 'DESUSO',
            'observacao': 'Foto e conferência salvas juntas',
        },
        content_type='multipart/form-data',
    )
    assert resposta.status_code == 200, resposta.get_json()
    item = client.get('/api/conferencia?setor=Setor%20A').get_json()['itens'][0]
    assert item['Foto']
    assert item['STATUS'] == 'OK'
    assert item['ESTADO'] == 'DESUSO'
    assert item['OBSERVACAO'] == 'Foto e conferência salvas juntas'
    assert resposta.get_json()['pasta'] == 'dados_trabalho\\Fotos'
    assert (FOTOS_DIR / item['Foto']).is_file()
    for nome_planilha in ('conferencia_resultado.xlsx', 'conferencia_campo.xlsx'):
        dados_salvos = pd.read_excel(DADOS_TRABALHO_DIR / nome_planilha).fillna('')
        assert dados_salvos.loc[0, 'Foto'] == item['Foto']
        assert dados_salvos.loc[0, 'ESTADO'] == 'DESUSO'
    assert client.get(f"/dados-web/fotos/{item['Foto']}").status_code == 200
    com_foto = client.get('/api/busca?foto=Com+foto')
    assert com_foto.get_json()['total'] == 1

    removida = post(
        client,
        '/api/itens/0/foto',
        data={'remover': 'sim'},
        content_type='multipart/form-data',
    )
    assert removida.status_code == 200, removida.get_json()
    sem_foto = client.get('/api/busca?foto=Sem+foto')
    assert sem_foto.get_json()['total'] == 2


def test_busca_filtra_itens_com_e_sem_plaqueta(client, monkeypatch):
    pagina = client.get('/').get_data(as_text=True)
    assert 'id="search-plaqueta"' in pagina
    assert '<option value="Sem plaqueta">Sem plaqueta</option>' in pagina

    dados = pd.DataFrame([
        {'Plaqueta': '', 'Descrição': 'Bem sem plaqueta', 'STATUS': 'DIVERG'},
        {'Plaqueta': None, 'Descrição': 'Outro sem plaqueta', 'STATUS': 'DIVERG'},
        {'Plaqueta': '000123', 'Descrição': 'Bem identificado', 'STATUS': 'OK'},
    ])
    monkeypatch.setattr(app_module, '_carregar_base', lambda: dados)

    sem_plaqueta = client.get('/api/busca?plaqueta=Sem+plaqueta').get_json()
    com_plaqueta = client.get('/api/busca?plaqueta=Com+plaqueta').get_json()

    assert sem_plaqueta['total'] == 2
    assert all(not item['Plaqueta'] for item in sem_plaqueta['itens'])
    assert com_plaqueta['total'] == 1
    assert com_plaqueta['itens'][0]['Plaqueta'] == '000123'


def test_busca_e_conferencia_paginam_resultados(client, monkeypatch):
    dados = pd.DataFrame([
        {
            'Plaqueta': str(indice),
            'Descrição': f'Bem {indice}',
            'Setor Atual': 'Setor A',
            'STATUS': 'PEND',
        }
        for indice in range(120)
    ])
    monkeypatch.setattr(app_module, '_carregar_base', lambda: dados)

    primeira_pagina = client.get('/api/busca').get_json()
    segunda_pagina = client.get('/api/busca?pagina=2').get_json()
    ultima_pagina = client.get('/api/busca?pagina=99').get_json()
    conferencia = client.get('/api/conferencia?setor=Setor+A&pagina=3').get_json()

    assert primeira_pagina['total'] == 120
    assert len(primeira_pagina['itens']) == 50
    assert primeira_pagina['pagina'] == 1
    assert len(segunda_pagina['itens']) == 50
    assert segunda_pagina['itens'][0]['Plaqueta'] == '50'
    assert ultima_pagina['pagina'] == 3
    assert len(ultima_pagina['itens']) == 20
    assert conferencia['total'] == 120
    assert len(conferencia['itens']) == 20

    invalida = client.get('/api/busca?por_pagina=101')
    assert invalida.status_code == 400


def test_lista_leilao_importa_csv_e_permite_vinculo_confirmado(client, monkeypatch):
    base = pd.DataFrame([{
        'Plaqueta': '000123',
        'Descrição': 'Mesa de reunião',
        'Setor Atual': 'Setor A',
        'STATUS': 'PEND',
        'ESTADO': 'EM USO',
    }], index=[7])
    monkeypatch.setattr(app_module, '_carregar_base', lambda: base)

    importacao = post(
        client,
        '/api/leilao/importar',
        data={
            'arquivo': (
                BytesIO('Descrição;Setor;Data\nMesa reunião;Setor A;2020-01-01\n'.encode()),
                'leilao.csv',
            ),
        },
        content_type='multipart/form-data',
    )
    assert importacao.status_code == 200, importacao.get_json()
    assert importacao.get_json()['importados'] == 1

    item = client.get('/api/leilao').get_json()['itens'][0]
    assert item['situacao'] == 'PENDENTE_IDENTIFICACAO'
    assert 'sugestoes' not in item
    sugestoes = client.get(f"/api/leilao/{item['id']}/sugestoes").get_json()
    assert sugestoes['sugestoes'][0]['plaqueta'] == '000123'
    assert base.iloc[0]['STATUS'] == 'PEND'

    vinculo = post(
        client,
        f"/api/leilao/{item['id']}/vincular",
        json={'plaqueta': '000123'},
    )
    assert vinculo.status_code == 200, vinculo.get_json()
    assert 'Nenhuma baixa foi realizada' in vinculo.get_json()['mensagem']
    atualizado = client.get('/api/leilao').get_json()['itens'][0]
    assert atualizado['situacao'] == 'IDENTIFICADO'
    assert atualizado['plaqueta_vinculada'] == '000123'
    assert base.iloc[0]['STATUS'] == 'PEND'


def test_lista_leilao_importa_excel_e_inclui_item_manual(client):
    workbook = BytesIO()
    pd.DataFrame([{'Bem': 'Cadeira', 'Plaqueta': '', 'Setor atual': 'Setor B'}]).to_excel(
        workbook, index=False
    )
    importacao = post(
        client,
        '/api/leilao/importar',
        data={'arquivo': (BytesIO(workbook.getvalue()), 'leilao.xlsx')},
        content_type='multipart/form-data',
    )
    assert importacao.status_code == 200, importacao.get_json()

    manual = post(
        client,
        '/api/leilao/manual',
        json={
            'descricao': 'Armário de aço',
            'setor': 'Setor C',
            'data': '2021-05-10',
        },
    )
    assert manual.status_code == 200, manual.get_json()
    itens = client.get('/api/leilao').get_json()['itens']
    assert len(itens) == 2
    assert {item['descricao'] for item in itens} == {'Cadeira', 'Armário de aço'}
    assert all(item['situacao'] == 'PENDENTE_IDENTIFICACAO' for item in itens)


def test_rota_entrega_miniatura_redimensionada(client):
    FOTOS_DIR.mkdir(parents=True, exist_ok=True)
    caminho = FOTOS_DIR / 'miniatura_teste.png'
    Image.new('RGB', (300, 150), color=(20, 40, 60)).save(caminho)

    resposta = client.get('/dados-web/fotos/miniatura_teste.png?miniatura=1')

    assert resposta.status_code == 200
    assert resposta.mimetype == 'image/jpeg'
    with Image.open(BytesIO(resposta.data)) as miniatura:
        assert miniatura.size == (96, 48)


def test_relatorio_gerencial_e_backup_web_restauravel(client):
    assert preparar_base(client).status_code == 200
    relatorio = post(client, '/api/relatorios/pendencias', json={})
    assert relatorio.status_code == 200, relatorio.get_json()
    nome_relatorio = relatorio.get_json()['arquivo']
    assert client.get(f'/api/relatorios/{nome_relatorio}').status_code == 200

    item = client.get('/api/conferencia?setor=Setor%20A').get_json()['itens'][0]
    assert post(
        client,
        f"/api/itens/{item['id']}/editar",
        json={'status': 'OK', 'estado': 'EM USO', 'observacao': 'Estado salvo'},
    ).status_code == 200
    backup = post(client, '/api/backups', json={})
    assert backup.status_code == 200
    nome_backup = backup.get_json()['arquivo']

    assert post(
        client,
        f"/api/itens/{item['id']}/editar",
        json={'status': 'BAIXA', 'estado': 'EM USO', 'observacao': 'Mudança posterior'},
    ).status_code == 200
    restauracao = post(
        client,
        f'/api/backups/{nome_backup}/restaurar',
        json={'confirmado': True},
    )
    assert restauracao.status_code == 200, restauracao.get_json()

    restaurado = client.get('/api/conferencia?setor=Setor%20A').get_json()['itens'][0]
    assert restaurado['STATUS'] == 'OK'
    assert restaurado['OBSERVACAO'] == 'Estado salvo'
    assert list(BACKUPS_DIR.glob('backup_web_*.zip'))


def test_relatorios_da_busca_exportam_exatamente_os_filtros_aplicados(client):
    assert preparar_base(client).status_code == 200
    filtros = {
        'termo': '',
        'setores': ['Setor A'],
        'status': ['PEND'],
        'foto': [],
        'data_inicial': '',
        'data_final': '',
    }

    excel = post(client, '/api/relatorios/busca/excel', json=filtros)
    assert excel.status_code == 200, excel.get_json()
    assert excel.get_json()['total'] == 1
    resposta_excel = client.get(
        f"/api/relatorios/{excel.get_json()['arquivo']}"
    )
    assert resposta_excel.status_code == 200
    workbook = load_workbook(BytesIO(resposta_excel.data), read_only=True)
    assert 'Itens_Filtrados' in workbook.sheetnames

    pdf = post(client, '/api/relatorios/busca/pdf', json=filtros)
    assert pdf.status_code == 200, pdf.get_json()
    resposta_pdf = client.get(f"/api/relatorios/{pdf.get_json()['arquivo']}")
    assert resposta_pdf.status_code == 200
    assert resposta_pdf.data.startswith(b'%PDF')

    vazio = post(
        client,
        '/api/relatorios/busca/excel',
        json={**filtros, 'termo': 'texto sem correspondência'},
    )
    assert vazio.status_code == 400


def test_restauracao_exige_confirmacao_e_backup_web_valido(client):
    assert preparar_base(client).status_code == 200
    backup = post(client, '/api/backups', json={}).get_json()['arquivo']

    sem_confirmacao = post(
        client,
        f'/api/backups/{backup}/restaurar',
        json={'confirmado': False},
    )
    assert sem_confirmacao.status_code == 400

    arquivo_desktop = BACKUPS_DIR / 'backup_web_invalido.zip'
    arquivo_desktop.write_bytes(b'nao e um zip')
    invalido = post(
        client,
        '/api/backups/backup_web_invalido.zip/restaurar',
        json={'confirmado': True},
    )
    assert invalido.status_code == 400
